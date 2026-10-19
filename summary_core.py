from collections import defaultdict
from datetime import date, datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def read_data(filepath):
    workbook = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
    try:
        data = []
        for row in workbook.active.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            redeemed_at = row[0]
            product_type = str(row[1]) if row[1] else ""
            if isinstance(redeemed_at, datetime):
                day = redeemed_at.date()
            elif isinstance(redeemed_at, date):
                day = redeemed_at
            else:
                day = datetime.strptime(str(redeemed_at)[:10], "%Y-%m-%d").date()
            data.append((
                day,
                "代金券" in product_type,
                float(row[2] or 0),
                float(row[3] or 0),
                float(row[4] or 0),
            ))
        return data
    finally:
        workbook.close()


def aggregate(data):
    result = defaultdict(lambda: [0, 0, 0])
    for day, is_coupon, received, service_fee, merchant_amount in data:
        totals = result[(day, is_coupon)]
        totals[0] += received
        totals[1] += service_fee
        totals[2] += merchant_amount
    return result


def collect_dates(*aggregates):
    days = sorted({key[0] for values in aggregates for key in values})
    if not days:
        raise ValueError("输入文件没有可汇总的有效数据")
    return days


def generate_report(lankao_path, minquan_path, output_path, progress=None):
    notify = progress or (lambda _value: None)
    notify(0.1)
    lankao = aggregate(read_data(lankao_path))
    notify(0.3)
    minquan = aggregate(read_data(minquan_path))
    days = collect_dates(lankao, minquan)
    notify(0.5)

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "汇总"
    header_font = Font(name="等线", size=11, bold=True)
    header_align = Alignment(horizontal="center", vertical="center")
    data_font = Font(name="等线", size=11)
    data_align = Alignment(horizontal="right", vertical="center")
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    green = PatternFill(start_color="FF92D050", end_color="FF92D050", fill_type="solid")
    yellow = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
    orange = PatternFill(start_color="FFFFC000", end_color="FFFFC000", fill_type="solid")

    def set_cell(row, column, value, font=None, align=None, fill=None, number_format=None):
        cell = sheet.cell(row=row, column=column, value=value)
        if font is not None:
            cell.font = font
        if align is not None:
            cell.alignment = align
        if fill is not None:
            cell.fill = fill
        cell.border = border
        if number_format:
            cell.number_format = number_format
        return cell

    sheet.merge_cells("B1:J1")
    set_cell(1, 2, "商品", header_font, header_align)
    sheet.merge_cells("K1:S1")
    set_cell(1, 11, "代金券", header_font, header_align)
    sheet.merge_cells("A2:A3")
    set_cell(2, 1, "日期", header_font, header_align)
    set_cell(3, 1, None)

    groups = [
        ("B2:D2", "合计", green), ("E2:G2", "兰考", yellow), ("H2:J2", "民权", orange),
        ("K2:M2", "合计", green), ("N2:P2", "兰考", yellow), ("Q2:S2", "民权", orange),
    ]
    for merge_range, label, fill in groups:
        sheet.merge_cells(merge_range)
        column = openpyxl.utils.cell.column_index_from_string(merge_range[0])
        set_cell(2, column, label, header_font, header_align, fill)

    fills = [green] * 3 + [yellow] * 3 + [orange] * 3
    for index, (label, fill) in enumerate(zip(["订单实收", "服务费", "商家应得"] * 6, fills * 2)):
        set_cell(3, 2 + index, label, header_font, header_align, fill)

    for row in range(1, 4):
        for column in range(1, 20):
            sheet.cell(row=row, column=column).border = border
    sheet.column_dimensions["A"].width = 8.3
    sheet.column_dimensions["B"].width = 16.3
    sheet.column_dimensions["K"].width = 16.3
    for column in "CDEFGHIJLMNOPQRS":
        sheet.column_dimensions[column].width = 13

    for index, day in enumerate(days):
        row = 4 + index
        set_cell(row, 1, datetime(day.year, day.month, day.day), header_font, header_align, number_format="M/D")
        lankao_goods = lankao.get((day, False), [0, 0, 0])
        minquan_goods = minquan.get((day, False), [0, 0, 0])
        lankao_coupon = lankao.get((day, True), [0, 0, 0])
        minquan_coupon = minquan.get((day, True), [0, 0, 0])
        values = (
            [lankao_goods[i] + minquan_goods[i] for i in range(3)] + lankao_goods + minquan_goods +
            [lankao_coupon[i] + minquan_coupon[i] for i in range(3)] + lankao_coupon + minquan_coupon
        )
        for value_index, value in enumerate(values):
            set_cell(row, 2 + value_index, round(value, 2), data_font, data_align, number_format="#,##0.00")
        notify(0.5 + 0.45 * (index + 1) / len(days))

    try:
        workbook.save(output_path)
    finally:
        workbook.close()
    notify(1.0)
    return {"days": len(days), "first_date": days[0], "last_date": days[-1], "output": str(output_path)}
