import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl import Workbook

from summary_core import aggregate, collect_dates, generate_report, read_data


class SummaryCoreTests(unittest.TestCase):
    def test_reads_and_aggregates_products_and_coupons(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["核销时间", "商品类型", "订单实收", "服务费", "商家应得"])
            sheet.append([datetime(2026, 10, 1, 10), "普通商品", 100, 5, 95])
            sheet.append([date(2026, 10, 1), "代金券", 200, 10, 190])
            workbook.save(path)
            workbook.close()

            result = aggregate(read_data(path))
            self.assertEqual(result[(date(2026, 10, 1), False)], [100, 5, 95])
            self.assertEqual(result[(date(2026, 10, 1), True)], [200, 10, 190])
            self.assertEqual(collect_dates(result), [date(2026, 10, 1)])

    def test_empty_inputs_are_rejected_before_report_generation(self):
        with self.assertRaisesRegex(ValueError, "没有可汇总"):
            collect_dates({}, {})

    def test_generate_report_creates_expected_totals(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            inputs = []
            for name, value in (("lankao", 100), ("minquan", 200)):
                path = root / f"{name}.xlsx"
                workbook = Workbook()
                sheet = workbook.active
                sheet.append(["核销时间", "商品类型", "订单实收", "服务费", "商家应得"])
                sheet.append([date(2026, 10, 1), "普通商品", value, 5, value - 5])
                workbook.save(path)
                workbook.close()
                inputs.append(path)
            output = root / "summary.xlsx"
            result = generate_report(inputs[0], inputs[1], output)
            self.assertEqual(result["days"], 1)
            workbook = openpyxl.load_workbook(output, data_only=True)
            self.assertEqual(workbook.active.cell(4, 2).value, 300)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
